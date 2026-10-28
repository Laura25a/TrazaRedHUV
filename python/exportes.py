"""
exportes.py — Arma los archivos descargables (CSV, Excel y PDF) del registro.

Una sola forma de exportar para los dos libros (radicación y contable) y para
el formato de cada remisión. Los colores del PDF son los mismos de la interfaz.
"""
import csv
import io
from datetime import datetime
from typing import Iterable, List, Optional, Sequence
from zoneinfo import ZoneInfo

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

PETROLEO = colors.HexColor("#0a4f4d")
PETROLEO_MEDIO = colors.HexColor("#0e6b68")
PAPEL = colors.HexColor("#fffdf8")
LINEA = colors.HexColor("#cfdcd8")
TINTA = colors.HexColor("#18302f")
CORAL = colors.HexColor("#e0634a")

ZONA = ZoneInfo("America/Bogota")   # todo lo descargable va en hora de Colombia

MIME = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "json": "application/json",
}


def archivo(contenido: bytes, nombre: str, tipo: str) -> Response:
    """Respuesta que el navegador descarga con ese nombre."""
    return Response(content=contenido, media_type=MIME[tipo],
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"',
                             "Cache-Control": "no-store"})


def texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.astimezone(ZONA).strftime("%d/%m/%Y %H:%M") if v.tzinfo else v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float):
        if v.is_integer():
            return f"{int(v):,}".replace(",", ".")
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


def a_csv(columnas: Sequence[str], filas: Iterable[Sequence]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")          # ; para que Excel en español lo abra bien
    w.writerow(columnas)
    for f in filas:
        w.writerow([texto(v) for v in f])
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM: tildes bien en Excel


def a_xlsx(hojas: List[dict]) -> bytes:
    """hojas: [{"titulo", "encabezado": [líneas], "columnas", "filas", "anchos"?}]"""
    wb = Workbook()
    wb.remove(wb.active)
    relleno = PatternFill("solid", fgColor="0A4F4D")
    borde = Border(bottom=Side(style="thin", color="CFDCD8"))
    for h in hojas:
        ws = wb.create_sheet(h["titulo"][:31])
        fila = 1
        for linea in h.get("encabezado", []):
            ws.cell(row=fila, column=1, value=linea).font = Font(bold=(fila == 1), size=13 if fila == 1 else 10,
                                                                 color="18302F")
            fila += 1
        if h.get("encabezado"):
            fila += 1
        for j, c in enumerate(h["columnas"], start=1):
            cel = ws.cell(row=fila, column=j, value=c)
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = relleno
            cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = ws.cell(row=fila + 1, column=1)
        inicio = fila
        for f in h["filas"]:
            fila += 1
            for j, v in enumerate(f, start=1):
                if isinstance(v, datetime) and v.tzinfo:
                    v = v.astimezone(ZONA).replace(tzinfo=None)
                cel = ws.cell(row=fila, column=j, value=v)
                cel.border = borde
                cel.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(v, datetime):
                    cel.number_format = "dd/mm/yyyy hh:mm"
        anchos = h.get("anchos") or [max(12, min(48, len(str(c)) + 6)) for c in h["columnas"]]
        for j, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(j)].width = a
        if h["filas"]:
            ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(h['columnas']))}{fila}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_estilos = getSampleStyleSheet()
E_TITULO = ParagraphStyle("t", parent=_estilos["Title"], fontName="Times-Bold", fontSize=18,
                          textColor=TINTA, spaceAfter=2, alignment=0)
E_SUB = ParagraphStyle("s", parent=_estilos["Normal"], fontSize=8.5, textColor=colors.HexColor("#3d5654"))
E_CELDA = ParagraphStyle("c", parent=_estilos["Normal"], fontSize=7.6, leading=9.4, textColor=TINTA)
E_CELDA_B = ParagraphStyle("cb", parent=E_CELDA, fontName="Helvetica-Bold")
E_SECCION = ParagraphStyle("sec", parent=_estilos["Normal"], fontName="Helvetica-Bold", fontSize=8.5,
                           textColor=colors.white)
E_PIE = ParagraphStyle("p", parent=_estilos["Normal"], fontSize=7, textColor=colors.HexColor("#3d5654"))
E_CENTRO = ParagraphStyle("cc", parent=E_CELDA, alignment=TA_CENTER)


def unidad(u) -> str:
    """UCUM → como se lee en papel."""
    return {"mm[Hg]": "mmHg", "Cel": "°C", "{score}": "/15", "/min": "/min", "%": "%"}.get(u or "", u or "")


def _p(v, estilo=E_CELDA):
    return Paragraph(texto(v).replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br/>"), estilo)


def _pie(lienzo, doc):
    lienzo.saveState()
    lienzo.setFont("Helvetica", 7)
    lienzo.setFillColor(colors.HexColor("#3d5654"))
    lienzo.drawString(doc.leftMargin, 0.8 * cm,
                      f"TrazaRed HUV · generado el {datetime.now(ZONA):%d/%m/%Y %H:%M} · documento derivado del registro, no editable")
    lienzo.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.8 * cm, f"Página {doc.page}")
    lienzo.restoreState()


def tabla_pdf(titulo: str, subtitulo: str, columnas: Sequence[str], filas: Iterable[Sequence],
              anchos: Optional[Sequence[float]] = None, apaisado=True, nota: Optional[str] = None) -> bytes:
    buf = io.BytesIO()
    pagina = landscape(A4) if apaisado else A4
    doc = SimpleDocTemplate(buf, pagesize=pagina, leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                            topMargin=1.3 * cm, bottomMargin=1.5 * cm, title=titulo)
    historia = [Paragraph("HOSPITAL UNIVERSITARIO DEL VALLE «EVARISTO GARCÍA» E.S.E.", E_SUB),
                Paragraph(titulo, E_TITULO), Paragraph(subtitulo, E_SUB), Spacer(1, 0.35 * cm)]
    datos = [[_p(c, ParagraphStyle("h", parent=E_CELDA_B, textColor=colors.white)) for c in columnas]]
    datos += [[_p(v) for v in f] for f in filas]
    ancho_util = pagina[0] - doc.leftMargin - doc.rightMargin
    if anchos:
        total = sum(anchos)
        anchos = [ancho_util * a / total for a in anchos]
    t = Table(datos, colWidths=anchos, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), PETROLEO),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINEA),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6faf8")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    historia.append(t)
    if nota:
        historia += [Spacer(1, 0.3 * cm), Paragraph(nota, E_PIE)]
    doc.build(historia, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


# ── Formato de una remisión (constancia) ─────────────────────────────────────
def _seccion(titulo: str, color=PETROLEO):
    t = Table([[Paragraph(titulo, E_SECCION)]], colWidths=["100%"])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), color),
                           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    return t


def _campos(pares: Sequence[tuple], columnas: int, ancho: float):
    celdas, fila = [], []
    for k, v in pares:
        fila.append([Paragraph(k.upper(), ParagraphStyle("k", parent=E_CELDA, fontSize=6.3,
                                                         textColor=colors.HexColor("#3d5654"))),
                     _p(v, E_CELDA_B)])
        if len(fila) == columnas:
            celdas.append(fila); fila = []
    if fila:
        celdas.append(fila + [""] * (columnas - len(fila)))
    t = Table(celdas, colWidths=[ancho / columnas] * columnas)
    t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, TINTA), ("INNERGRID", (0, 0), (-1, -1), 0.4, LINEA),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def constancia_pdf(f: dict) -> bytes:
    """f = lo que devuelve registro.formato_remision(): arma el formato TR-FR-01."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.3 * cm, bottomMargin=1.6 * cm, title=f["remision"]["radicado"])
    W = A4[0] - doc.leftMargin - doc.rightMargin
    r, p = f["remision"], f["paciente"]
    h = []
    enc = Table([[Paragraph("<b>HUV</b>", ParagraphStyle("l", parent=E_TITULO, textColor=PETROLEO_MEDIO,
                                                          alignment=TA_CENTER)),
                  [Paragraph("Hospital Universitario del Valle «Evaristo García» E.S.E.", E_CENTRO),
                   Paragraph("<b>REGISTRO DE REFERENCIA Y CONTRARREFERENCIA DE PACIENTES</b>",
                             ParagraphStyle("tt", parent=E_CENTRO, fontSize=10.5, leading=13))],
                  [Paragraph("Código: TR-FR-01", E_CELDA), Paragraph("Versión: 02", E_CELDA),
                   Paragraph(f"Folio {r['folio']:04d} · Asiento {r['id']:06d}", E_CELDA)]]],
                colWidths=[2.6 * cm, W - 7.6 * cm, 5 * cm])
    enc.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1.2, TINTA), ("INNERGRID", (0, 0), (-1, -1), 1, TINTA),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    h += [enc, Spacer(1, 0.3 * cm)]
    h.append(_campos([("N.° de radicado", r["radicado"]), ("Radicada", r.get("radicada_en")),
                      ("Estado", f["estado_texto"]), ("Tiempo", f["tiempo_texto"])], 4, W))
    h += [Spacer(1, 0.25 * cm), _seccion("1. IDENTIFICACIÓN DEL PACIENTE"),
          _campos([("Nombres y apellidos", p.get("nombre")), ("Documento", p.get("documento")),
                   ("Edad / sexo", f"{p.get('edad') or '—'} años · {p.get('genero') or '—'}"), ("EPS", p.get("eps")),
                   ("Teléfono", p.get("telefono") or "—"), ("Tipo de sangre", p.get("tipo_sangre") or "—"),
                   ("Alergias", p.get("alergias") or "—"), ("", "")], 4, W),
          Spacer(1, 0.25 * cm), _seccion("2. DATOS DE LA SOLICITUD"),
          _campos([("Institución que remite", r.get("institucion_origen")),
                   ("Institución receptora", r.get("institucion_destino")),
                   ("Médico que radica", r.get("asento_nombre") or "—"),
                   ("Convenio con destino", "Vigente" if r.get("convenio_vigente") else "Sin convenio"),
                   ("Diagnóstico", f["diagnostico"]), ("Motivo de remisión", f["requiere"])], 2, W)]
    if f.get("signos"):
        h.append(_campos([(s["tipo"], f"{texto(s['valor'])} {unidad(s.get('unidad'))}") for s in f["signos"]], 4, W))
    h += [Spacer(1, 0.25 * cm), _seccion("3. BITÁCORA DE GESTIÓN · cada fila es un registro firmado, no editable")]
    filas = [["N.°", "Fecha / hora", "Institución · medio", "Respuesta anotada", "Causa", "Registró"]]
    for e in f["bitacora"]:
        filas.append([f"{e['orden']:02d}", texto(e.get("fecha")), f"{e.get('institucion') or '—'} · {e.get('medio') or ''}",
                      e.get("respuesta") or "", e.get("causa_etiqueta") or "", e.get("por") or ""])
    if len(filas) == 1:
        filas.append(["—", "", "", "Sin gestiones anotadas", "", ""])
    t = Table([[_p(c, E_CELDA_B if i == 0 else E_CELDA) for c in fila] for i, fila in enumerate(filas)],
              colWidths=[0.9 * cm, 2.5 * cm, 3.8 * cm, W - 12.6 * cm, 3 * cm, 2.4 * cm], repeatRows=1)
    t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, TINTA), ("INNERGRID", (0, 0), (-1, -1), 0.4, LINEA),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef4f2")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    h.append(t)
    ia = f.get("pronostico")
    if ia:
        h += [Spacer(1, 0.25 * cm), _seccion("4. APOYO DE IA · orientativo, no reemplaza el criterio del equipo",
                                             colors.HexColor("#2f6fb3")),
              _campos([("Probabilidad de que termine efectiva", f"{round(100 * ia['probabilidad_efectiva'])} %"),
                       ("Al radicar", f"{round(100 * (ia.get('probabilidad_al_radicar') or 0))} %"),
                       ("Riesgo de estancamiento", ia.get("riesgo_estancamiento") or ia.get("riesgo") or "—"),
                       ("Modelo", ia.get("version_modelo") or "—")], 4, W)]
    costos = f.get("costos")
    if costos:
        h += [Spacer(1, 0.25 * cm), _seccion("5. COSTOS ASOCIADOS AL RADICADO")]
        filas = [["Concepto", "Cantidad", "Valor", "Cobro a"]]
        filas += [[c["concepto"], f"{texto(c['cantidad'])} {c['unidad']}", f"$ {c['valor']:,.0f}".replace(",", "."),
                   costos["cobro_a"]] for c in costos["conceptos"]]
        filas.append(["Total", "", f"$ {costos['total']:,.0f}".replace(",", "."), costos["estado_cobro_texto"]])
        t = Table([[_p(c, E_CELDA_B if i in (0, len(filas) - 1) else E_CELDA) for c in fila] for i, fila in enumerate(filas)],
                  colWidths=[W - 9 * cm, 2.6 * cm, 3 * cm, 3.4 * cm])
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, TINTA), ("INNERGRID", (0, 0), (-1, -1), 0.4, LINEA),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef4f2"))]))
        h.append(t)
    c = f["cierre"]
    h += [Spacer(1, 0.25 * cm), _seccion("6. CIERRE DE LA REMISIÓN"),
          _campos([("Institución que acepta", c.get("institucion_acepta") or "—"),
                   ("Fecha / hora de cierre", c.get("cerrada_en") or "—"),
                   ("Resultado", c.get("resultado_texto") or "—")], 3, W)]
    firmas = Table([[Paragraph(f"<b>{r.get('asento_nombre') or '—'}</b><br/>Médico que radica", E_CELDA),
                     Paragraph(f"<b>{f.get('responsable') or '—'}</b><br/>Oficina de Referencia · responsable de la gestión", E_CELDA)]],
                   colWidths=[W / 2 - 0.5 * cm, W / 2 - 0.5 * cm], spaceBefore=26)
    firmas.setStyle(TableStyle([("LINEABOVE", (0, 0), (0, 0), 0.8, TINTA), ("LINEABOVE", (1, 0), (1, 0), 0.8, TINTA)]))
    h += [KeepTogether([firmas]), Spacer(1, 0.4 * cm),
          Paragraph("Documento generado por TrazaRed HUV a partir de los registros del sistema (PostgreSQL, MongoDB y HAPI FHIR). "
                    "La bitácora no puede modificarse; las correcciones se anotan como nuevas entradas con referencia a la original.", E_PIE)]
    doc.build(h, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()
